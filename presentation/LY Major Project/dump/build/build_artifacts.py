from pathlib import Path
from copy import deepcopy
import json, zipfile, re, posixpath, hashlib
from lxml import etree as E
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[4]
OUT=ROOT/'presentation/LY Major Project'
BUILD=Path(__file__).resolve().parent
ASSETS=BUILD/'assets';ASSETS.mkdir(exist_ok=True)
M=json.loads((ROOT/'models/metrics.json').read_text())
RESULTS=json.loads((ROOT/'report/data/results.json').read_text())
FONT='Times New Roman'
TITLE='Smart Moisture Advisor'
FULL='Smart Moisture Advisor An AI Driven Irrigation Advisory System for Sugarcane'
REFS=[
 '[1] R. G. Allen, L. S. Pereira, D. Raes, and M. Smith, Crop Evapotranspiration: Guidelines for Computing Crop Water Requirements, FAO Irrigation and Drainage Paper 56. Rome: FAO, 1998. https://www.fao.org/4/x0490e/x0490e00.htm',
 '[2] J. Doorenbos and A. H. Kassam, Yield Response to Water, FAO Irrigation and Drainage Paper 33. Rome: FAO, 1979. Bibliographic record: https://www.fao.org/4/y3655e/y3655e03.htm',
 '[3] C. Kamienski et al., "Smart Water Management Platform: IoT-Based Precision Irrigation for Agriculture," Sensors, vol. 19, no. 2, article 276, 2019. doi: 10.3390/s19020276. https://www.mdpi.com/1424-8220/19/2/276',
 '[4] L. Breiman, "Random forests," Machine Learning, vol. 45, pp. 5-32, 2001. doi: 10.1023/A:1010933404324. https://doi.org/10.1023/A:1010933404324',
 '[5] J. H. Friedman, "Greedy function approximation: A gradient boosting machine," The Annals of Statistics, vol. 29, no. 5, pp. 1189-1232, 2001. doi: 10.1214/aos/1013203451. https://doi.org/10.1214/aos/1013203451',
 '[6] F. Pedregosa et al., "Scikit-learn: Machine Learning in Python," Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011. https://www.jmlr.org/papers/v12/pedregosa11a.html',
 '[7] Scikit-learn developers, "GroupShuffleSplit," official API documentation. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GroupShuffleSplit.html',
 '[8] Express contributors, "Express 5.x API," official documentation. https://expressjs.com/en/5x/api.html',
 '[9] Leaflet contributors, "Leaflet API reference," official documentation. https://leafletjs.com/reference.html',
 '[10] K. Verma and K. Soni, Smart Moisture Advisor, project implementation and repository artifacts, Interior-Gardener/SMA, local main checkout reviewed 19 October 2026. https://github.com/Interior-Gardener/SMA. Evidence: ml/train.py, ml/agronomy.py, src/, models/metrics.json, models/test_fixtures.json, report/data/results.json, report/data/extra.json, and presentation/screenshots/.',
]

def font(size,bold=False):
 return ImageFont.truetype('C:/Windows/Fonts/timesbd.ttf' if bold else 'C:/Windows/Fonts/times.ttf',size)
def wrap(draw,s,f,w):
 lines=[]
 for para in s.split('\n'):
  cur=''
  for word in para.split():
   test=(cur+' '+word).strip()
   if draw.textlength(test,font=f)>w and cur:lines.append(cur);cur=word
   else:cur=test
  lines.append(cur)
 return lines
def drawtext(d,box,text,size=31,bold=False,fill='#17221e',align='center'):
 x,y,w,h=box;f=font(size,bold);ls=wrap(d,text,f,w-24);lh=size*1.17
 yy=y+(h-len(ls)*lh)/2
 for line in ls:
  xx=x+(w-d.textlength(line,font=f))/2 if align=='center' else x+12
  d.text((xx,yy),line,font=f,fill=fill);yy+=lh

def architecture():
 im=Image.new('RGB',(2000,1080),'white');d=ImageDraw.Draw(im)
 d.text((25,20),'Offline model development',font=font(43,True),fill='#000')
 boxes=[(25,100,330,165,'Plot CSV\n1,000 geofenced plots'),(425,100,330,165,'Clean and expand\n12 scenarios per plot'),(825,100,330,165,'Agronomy labels\nFAO-56 and FAO-33'),(1225,100,330,165,'Train and evaluate\nscikit-learn ensembles'),(1625,100,350,165,'Export artifacts\n7 tree models and zones')]
 def box(x,y,w,h,t):d.rectangle((x,y,x+w,y+h),fill='#F3F5F4',outline='#285447',width=3);drawtext(d,(x,y,w,h),t)
 def arrow(x1,y1,x2,y2):
  d.line((x1,y1,x2,y2),fill='#315B4D',width=5)
  if y1==y2:d.polygon([(x2,y2),(x2-17,y2-10),(x2-17,y2+10)],fill='#315B4D')
  else:d.polygon([(x2,y2),(x2-10,y2-17),(x2+10,y2-17)],fill='#315B4D')
 for b in boxes:box(*b)
 for i in range(4):arrow(boxes[i][0]+boxes[i][2]+6,182,boxes[i+1][0]-8,182)
 d.text((25,345),'Online advisory application',font=font(43,True),fill='#000')
 for b in [(25,445,330,190,'CSV farm loader\nNDVI imputation\nGeometry and quality flags'),(425,445,420,190,'Express REST API\nScenario validation\nPlot and simulator inputs'),(935,445,420,190,'JavaScript inference\nJSON tree traversal\nFour management zones'),(1445,445,530,190,'Decision engine\nRain adjustment and hydraulics\nFertigation and explanations')]:box(*b)
 arrow(355,540,415,540);arrow(845,540,925,540);arrow(1355,540,1435,540)
 d.line((1800,265,1800,325,1145,325,1145,435),fill='#315B4D',width=5);d.polygon([(1145,435),(1135,418),(1155,418)],fill='#315B4D')
 for b in [(425,795,590,205,'Browser dashboard\nMap, advisor, simulator, insights\nFour-language keyword chatbot'),(1150,795,420,205,'Pump scheduler\nWSPT list heuristic\nPower windows and capacity'),(1650,795,325,205,'Feedback log\nAccept or override\nLocal JSONL file')]:box(*b)
 d.line((1710,635,1710,700,720,700,720,785),fill='#315B4D',width=5);d.polygon([(720,785),(710,768),(730,768)],fill='#315B4D')
 d.line((1710,700,1360,700,1360,785),fill='#315B4D',width=5);d.polygon([(1360,785),(1350,768),(1370,768)],fill='#315B4D')
 arrow(1015,900,1140,900)
 d.line((720,1000,720,1040,1810,1040,1810,1010),fill='#315B4D',width=5);d.polygon([(1810,1010),(1800,1027),(1820,1027)],fill='#315B4D')
 im.save(ASSETS/'architecture.png')

def crops():
 # Crops retain the original UI pixels. They are never generated product screens.
 spec={
 'overview':('01_overview.png',(740,270,1425,465)),
 'map':('02b_map_zoom.png',(395,385,2085,1320)),
 'advisor':('03_advisor.png',(405,620,1520,1245)),
 'rain':('05_rain.png',(395,270,1080,650)),
 'simulator':('06_simulator.png',(1560,355,2115,802)),
 'scheduler':('07_scheduler.png',(395,715,2110,1325)),
 'chat':('09_chat.png',(420,570,2080,1150)),
 'kannada':('04_advisor_kannada.png',(405,625,1510,1210)),
 'fertigation':('03_advisor.png',(405,1760,1240,2275)),
 'explanation':('03_advisor.png',(1280,1760,2120,2180)),
 'curves':('03_advisor.png',(410,1270,2120,1740)),
 'insights':('08_insights.png',(400,515,2110,1070)),
 }
 for name,(f,box) in spec.items():Image.open(ROOT/'presentation/screenshots'/f).crop(box).save(ASSETS/(name+'.png'))
 (BUILD/'image_provenance.json').write_text(json.dumps({k:{'source':'presentation/screenshots/'+v[0],'crop_pixels':v[1]} for k,v in spec.items()},indent=2))

def p(doc,text,style=None,size=None,bold=False,after=7):
 q=doc.add_paragraph(style=style);q.paragraph_format.space_after=Pt(after)
 r=q.add_run(text);r.bold=bold or (style in ['Heading 1','Heading 2','Heading 3','Title']);r.font.name=FONT;r.font.color.rgb=RGBColor(0,0,0)
 if size:r.font.size=Pt(size)
 return q
def h(doc,text):return p(doc,text,'Heading 2',after=5)
def table(doc,headers,rows,widths=None,size=10):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,s in enumerate(headers):t.rows[0].cells[i].text=s
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=str(s)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr()
  cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tc=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   tc.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9ECEB' if ri==0 else ('F7F8F8' if ri%2==0 else 'FFFFFF'));tc.append(sh)
   mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');mar.append(x)
   tc.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.02
    for r in para.runs:r.font.size=Pt(size);r.bold=ri==0
 return t
def image(doc,name,width=6.75):
 q=doc.add_paragraph();q.paragraph_format.space_after=Pt(4);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.add_run().add_picture(str(ASSETS/(name+'.png')),width=Inches(width))
def page(doc,num,title):
 if num>1:doc.add_page_break()
 p(doc,title,'Heading 1',after=10)
def panel(doc,name,title,caption,w=3.28,hgt=1.50):
 im=Image.open(ASSETS/(name+'.png'));iw,ih=im.size
 # Fixed image envelope with aspect preservation and no distortion.
 scale=min(w/iw,hgt/ih);ww=iw*scale;hh=ih*scale
 q=doc.add_paragraph();q.paragraph_format.space_after=Pt(3);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
 q.add_run().add_picture(str(ASSETS/(name+'.png')),width=Inches(ww),height=Inches(hh))
 p(doc,title,size=10,bold=True,after=3)
 p(doc,caption,size=9.5,after=8)

def grid_panels(doc,items,height=1.45):
 # A borderless layout table groups paired figures and their captions.
 t=doc.add_table(rows=0,cols=2);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for idx in range(0,len(items),2):
  cells=t.add_row().cells
  for ci,item in enumerate(items[idx:idx+2]):
   c=cells[ci];c.width=Inches(3.43)
   c._tc.get_or_add_tcPr().append(OxmlElement('w:tcBorders'))
   first=c.paragraphs[0];first.paragraph_format.space_after=Pt(3);first.alignment=WD_ALIGN_PARAGRAPH.CENTER
   name,title,caption=item;im=Image.open(ASSETS/(name+'.png'));ww,hh=im.size;s=min(3.23/ww,height/hh)
   first.add_run().add_picture(str(ASSETS/(name+'.png')),width=Inches(ww*s),height=Inches(hh*s))
   q=c.add_paragraph();q.paragraph_format.space_after=Pt(3);r=q.add_run(title);r.bold=True;r.font.size=Pt(10)
   q=c.add_paragraph();q.paragraph_format.space_after=Pt(10);q.paragraph_format.line_spacing=1.0;r=q.add_run(caption);r.font.size=Pt(9.5)
 return t

def report():
 doc=Document();s=doc.sections[0];s.page_width=Inches(8.2677);s.page_height=Inches(11.6929)
 s.top_margin=Inches(.62);s.bottom_margin=Inches(.62);s.left_margin=Inches(.68);s.right_margin=Inches(.68);s.header_distance=Inches(.25);s.footer_distance=Inches(.28)
 for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
  st=doc.styles[name];st.font.name=FONT;st.font.color.rgb=RGBColor(0,0,0)
  st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(7)
 doc.styles['Normal'].font.size=Pt(12);doc.styles['Normal'].paragraph_format.line_spacing=1.16
 doc.styles['Title'].font.size=Pt(18);doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 1'].font.bold=True;doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].font.bold=True
 footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 rr=footer.add_run('Smart Moisture Advisor   |   ');rr.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.core_properties.author='Kartik Verma and Kushal Soni';doc.core_properties.title=FULL;doc.core_properties.subject='Final year major project report';doc.core_properties.keywords='SMA, irrigation advisory, sugarcane, agronomy-generated labels'
 p(doc,FULL,'Title',after=8)
 p(doc,'Major Project Report   2026-27\nDepartment of Computer Engineering\nK. J. Somaiya Institute of Technology',size=10.5,after=8)
 p(doc,'Students: Kartik Verma and Kushal Soni\nProject Guide: Dr. Nilesh Yadav',size=11,after=12)
 page(doc,1,'1  Introduction')
 p(doc,'Smart Moisture Advisor (SMA) is a browser-based decision-support prototype for sugarcane irrigation. It combines plot attributes with user-selected weather scenarios to estimate when irrigation is due, the water volume and pump duration required, and the consequences of delay. The implemented system covers 1,000 geofenced plots in Mandya district, Karnataka, across six taluks and ten villages [10].')
 p(doc,'Irrigation planning requires a balance between crop demand, available soil water, expected rainfall, and the practical availability of pumps and electricity. FAO-56 supplies the crop-coefficient and root-zone water-balance framework used to reason about these quantities [1]. SMA connects this agronomic calculation to a plot-level interface that also supports fertiliser advice, risk estimates, explanations, and scheduling.')
 h(doc,'Project approach')
 p(doc,'The supplied dataset contains vegetation indices, soil properties, seasonal climate, and plot geometry, but no recorded irrigation dates, irrigation volumes, harvest yields, or disease outcomes. The training pipeline therefore expands each plot into twelve seasonal scenarios and generates target values from agronomic equations with added noise. Seven tree-ensemble models learn these relationships, while K-Means partitions the observed plots into four management zones [10].')
 p(doc,'Python and scikit-learn handle offline training. Exported JSON trees support inference directly in Node.js, so the deployed web application does not need a Python prediction service. Express serves the API and static client. Leaflet maps the plots, and Chart.js displays predictions and evaluation outputs. English, Kannada, Hindi, and Marathi advisories share the same structured predictions [6], [8]-[10].')
 h(doc,'Principal outcome and interpretation')
 p(doc,'The saved held-out-plot evaluation records an irrigation-timing R-squared of 0.9646 and a mean absolute error of 0.6942 days. These results quantify agreement with generated agronomic labels. They do not establish field accuracy or measured water savings. The project demonstrates an integrated software workflow whose next requirement is validation against independent irrigation and harvest records [10].')

 page(doc,2,'2  Literature Survey and Research Gaps')
 h(doc,'Agronomic water-balance methods')
 p(doc,'Allen et al. describe reference evapotranspiration, crop coefficients, available root-zone water, and soil-water depletion in FAO-56 [1]. This provides a physical basis for deciding when the crop approaches a water-stress threshold. The framework remains dependent on appropriate weather, crop, and soil parameters. SMA adopts its crop-water and depletion concepts while using simplified parameter estimates suited to the available inputs. These simplifications require local calibration before operational use.')
 p(doc,'Doorenbos and Kassam relate relative yield reduction to relative evapotranspiration deficit using a crop yield-response factor [2]. Their model motivates the project\'s simulated loss from delayed irrigation. A water-response relationship alone does not identify a farmer\'s actual yield or allocate competing pumps. SMA adds a delay input and scheduling workflow, but its yield and delay-loss labels still inherit the assumptions of its agronomic generator.')
 h(doc,'Connected precision irrigation platforms')
 p(doc,'Kamienski et al. present SWAMP, an IoT platform for precision irrigation with multiple agricultural pilots and a performance study of its FIWARE infrastructure [3]. The paper demonstrates the engineering work required to connect sensing, water management, and scalable services. It also reports that scalability depends on platform configuration. SMA takes a smaller local architecture: a CSV-backed API, portable models, and a browser interface. This reduces prototype infrastructure requirements, while leaving live sensing and operational integration for future work.')
 p(doc,'The comparison is architectural rather than experimental. SWAMP reports deployments of an IoT platform, whereas SMA currently uses a static plot dataset and manually selected forecast scenarios. No shared benchmark or head-to-head experiment exists in the repository, so a superiority claim would be unsupported [3], [10].')
 h(doc,'Tree ensembles for tabular prediction')
 p(doc,'Breiman\'s Random Forest combines randomized tree predictors through aggregation [4]. SMA uses forests for irrigation timing and the two ordered risk classifications. The spread of individual irrigation trees provides a descriptive uncertainty signal. It is not a calibrated confidence interval. The model\'s performance depends on how representative the generated scenarios are of the conditions encountered during use.')
 p(doc,'Friedman\'s gradient boosting fits an additive model through successive improvements to a loss function [5]. SMA uses gradient-boosted trees for water requirement, irrigation depth, cane yield, and delay-related yield loss. These methods suit nonlinear relationships in numerical inputs. Their flexibility does not remove the need for independent targets, appropriate data splitting, and deployment monitoring.')
 h(doc,'Research implication')
 p(doc,'The literature supplies the physical equations, ensemble methods, and examples of irrigation platforms. SMA\'s contribution is their implementation in one plot-advisory workflow. The remaining research gap is evidence that this combined workflow produces reliable decisions and benefits under actual field conditions [1]-[5], [10].')

 page(doc,3,'3  Literature Survey and Research Gaps')
 h(doc,'Evaluation and reproducible implementation')
 p(doc,'Pedregosa et al. describe scikit-learn as a consistent Python interface for supervised and unsupervised learning [6]. SMA uses it to train the ensembles and management zones. The GroupShuffleSplit API separates groups rather than individual samples [7]. In this project the group is the plot identifier, so seasonal scenarios from a test plot do not appear in the training partition. This supports a more meaningful test than randomly splitting all scenario rows.')
 p(doc,'The split still uses the same scenario and label generator on both sides. In addition, village-median NDVI imputation precedes the group split in ml/train.py. Consequently, the evaluation should be described as a grouped prototype benchmark rather than a fully isolated assessment of every preprocessing operation. Fitting preprocessing only on training plots would strengthen a future experiment [10].')
 h(doc,'Gaps and the implemented response')
 table(doc,['Gap or constraint','SMA implementation','Remaining evidence needed'],[
  ['No observed target outcomes','FAO-based labels plus noisy seasonal scenarios','Independent irrigation, yield, and disease labels'],
  ['Advice spans separate operational tasks','One API links prediction, hydraulics, fertigation, and scheduling','Farmer usability and agronomist review'],
  ['Electricity windows constrain execution','Parallel-lane WSPT scheduling with a first-come comparison','Measured feeder and pump behaviour'],
  ['Users need reasons and local language','Median-substitution explanations and four language templates','Explanation and translation validation'],
 ],[1.65,2.55,2.65],10)
 p(doc,'The table identifies gaps relative to the project\'s operational requirements. It does not assert that every cited platform lacks these capabilities. The source code supports the listed SMA response, while the final column describes work that the present implementation has not established [10].',size=10.5,after=10)
 h(doc,'Uncertainty and generalisation')
 p(doc,'The saved additional analysis reports 42.7% coverage within one forest standard deviation and 77.4% within two standard deviations for irrigation timing. These values do not support interpreting the displayed spread as a conventional 68% or 95% interval. The six taluks also have different held-out sample counts and errors, so a pooled score cannot establish uniform reliability across locations [10].')
 p(doc,'Crop age is a stable hash-derived estimate when no user override is supplied. Forecast rain is a scenario input, and the chatbot detects keywords rather than using a language model. The project has no demonstrated live satellite ingestion, weather-feed connector, pump control, or agronomic diagnosis service. Clear boundaries are necessary when interpreting the integrated interface [10].')

 page(doc,4,'4  Problem Definition Objectives and Scope')
 h(doc,'Problem definition')
 p(doc,'Develop a plot-specific sugarcane advisory system that converts available vegetation, soil, and climate attributes into understandable irrigation guidance, then helps allocate limited pumping capacity. The system must expose relevant assumptions and allow a person to review the recommendation. Its current dataset and generated labels constrain the claims that can be made about agronomic performance [10].')
 h(doc,'Objectives')
 for t in [
  'Prepare a usable plot dataset by parsing geofences, imputing missing NDVI, and flagging unusual values within villages.',
  'Estimate irrigation timing, crop water requirement, net irrigation depth, yield, delay-related yield loss, water stress, and disease or pest risk.',
  'Convert model outputs into method-specific gross water volumes, pump hours, rainfall-adjusted advice, and stage-wise fertigation quantities.',
  'Provide maps, input simulation, model diagnostics, and advisory text in English, Kannada, Hindi, and Marathi.',
  'Schedule irrigation jobs within power windows and feeder capacity, compare the heuristic with first-come rotation, and support accept or override feedback.',
 ]:p(doc,t,'List Bullet',after=6)
 h(doc,'Functional scope')
 p(doc,'The implemented application supports 1,000 CSV-backed plots, four irrigation methods, seven predictive models, and four management zones. The Farm Advisor uses a plot ID and optionally a revised crop age. The simulator accepts bounded numerical inputs, while the scheduler takes geographic filters, power windows, pump capacity, and a planning horizon. Feedback appends to a local JSONL file [10].')
 h(doc,'Operating assumptions')
 p(doc,'Irrigation efficiencies are configured as 0.55 for flood, 0.65 for furrow, 0.75 for sprinkler, and 0.90 for drip. Pump flow defaults to 50 cubic metres per hour. Water and energy comparisons use 70 mm of gross irrigation every eight days and a 3.73 kW pump. These are application parameters rather than measurements of each farm\'s equipment [10].')
 h(doc,'Scope boundaries and acceptance')
 p(doc,'The prototype provides decision support without actuating pumps. Sensor procurement, live forecast ingestion, authenticated multi-user operation, production databases, and measured field outcomes remain outside the implemented scope. Acceptance for this submission concerns an operational local application, traceable model artifacts, consistent Python-to-JavaScript predictions, and readable outputs. It does not imply agronomic deployment readiness.')
 p(doc,'The existing 23 tests pass in the reviewed checkout. They cover API behaviour, scenario effects, multilingual advisory generation, scheduler constraints, and inference parity with saved Python fixtures. Passing software tests demonstrates the tested behaviour and cannot replace a controlled field trial [10].')

 page(doc,5,'5  System Architecture and Data Flow')
 image(doc,'architecture',6.85)
 p(doc,'Figure 1. Architecture derived from ml/train.py, src/app.js, src/data/farms.js, src/ml/, and src/engine/. Model training is offline. The running Node.js service reads the CSV and exported models; it does not retrain during requests [10].',size=9.5,after=10)
 h(doc,'Offline preparation and training')
 p(doc,'The Python pipeline renames the dataset columns, fills six missing NDVI values using village medians, and creates 12,000 weather and crop-age scenarios. Agronomic functions generate labels with stochastic noise. An 80/20 GroupShuffleSplit separates 800 training plots from 200 test plots, producing 9,600 and 2,400 scenario rows. The yield-loss model uses three delay replications, giving 36,000 rows in total [7], [10].')
 h(doc,'Online prediction and delivery')
 p(doc,'At startup the farm loader parses plot polygons, computes centroids, estimates crop age, and records data-quality flags. The model registry loads JSON trees and zone centroids. Express validates API inputs, applies scenario settings, and calls the advisory engine. The browser requests summaries, plot details, simulations, schedules, or chat answers and renders the returned data [8]-[10].')
 h(doc,'Advice and feedback path')
 p(doc,'The engine combines predictions with irrigation efficiency, pump flow, and growth stage. Separate modules provide feature-sensitivity explanations, language templates, and pump scheduling. Accepting or overriding an advisory records a local feedback entry. The current training pipeline does not automatically consume this log, so feedback collection and automatic model improvement are separate capabilities [10].')

 page(doc,6,'6  Implementation and Important Algorithms')
 h(doc,'Models and targets')
 table(doc,['Target','Implemented model','Output'],[
  ['Next irrigation','Random Forest','Days and tree spread'],
  ['Crop water requirement','Gradient Boosting','ETc in mm/day'],
  ['Net irrigation depth','Gradient Boosting','Millimetres'],
  ['Cane yield','Gradient Boosting','Tonnes per hectare'],
  ['Loss from irrigation delay','Gradient Boosting','Percentage loss'],
  ['Water stress','Random Forest classifier','Four ordered classes'],
  ['Disease and pest risk','Random Forest classifier','Three ordered classes'],
 ],[2.25,2.35,2.25],10)
 p(doc,'A separate K-Means model assigns four management zones from standardised observed features. It is unsupervised and has no accuracy value comparable to the seven supervised targets [10].',size=10.5,after=9)
 h(doc,'Agronomic computation')
 p(doc,'The label generator estimates reference evapotranspiration with a Hargreaves-style equation and an assumed humidity-dependent temperature range. It blends an age-based sugarcane coefficient with an NDVI-based coefficient to obtain ETc = Kc x ET0. Root-zone depletion is Dr = 1000 x (FC - soil moisture) x root depth. Here FC is field capacity and RAW is readily available water. Days to irrigation equal (RAW - Dr + effective rain) / ETc, bounded to 0-15 days [1], [10].')
 p(doc,'Effective forecast rain is max(0, 0.8 x (rain - 3)). Net irrigation depth divided by method efficiency gives gross depth. Volume in cubic metres equals gross millimetres x area in hectares x 10; dividing by pump flow gives operating hours. The FAO-33 yield-response concept supports the simulated loss curve, supplemented by project-specific stage and heat factors [2], [10].')
 h(doc,'Inference and explanation')
 p(doc,'The JavaScript evaluator traverses exported tree arrays and casts numerical features to float32 to match scikit-learn. Ordered classifiers use the first class whose cumulative probability reaches 0.5. Explanations replace one input with its regional median and report the resulting change in irrigation timing. These are sensitivity effects and should not be described as SHAP values or causal effects [10].')
 h(doc,'Scheduling and workflow')
 p(doc,'The scheduler orders jobs by weight divided by pumping duration. Weight combines urgency, predicted tonnes at risk, and stress probability. It packs jobs into parallel feeder lanes, pauses at power-window boundaries, and compares the schedule with first-come rotation. This is a heuristic for the implemented constrained problem, with no global optimality guarantee. Node.js serves the saved models through REST endpoints; core application logic requires no change to produce the report or demonstration [8], [10].')

 page(doc,7,'7  Application Outputs and Scenario Results')
 p(doc,'The following are crops of existing application screenshots. Values describe the selected summer-dry or rainfall scenario, rather than observed farm outcomes. Full screens and saved scenario records are retained in the repository [10].',size=10,after=9)
 grid_panels(doc,[
  ('overview','Figure 2  Regional dry-spell overview','The dry-spell screen reports 637 plots due today and 204 with moderate or severe stress. Regional counts summarise the selected seven-day dry and +8 degree temperature scenario.'),
  ('map','Figure 3  Plot map and management layers','The Maddur map displays geofenced plots coloured by water-stress class. Layer choices also expose irrigation timing, vegetation, moisture, risk, and management zones.'),
  ('advisor','Figure 4  Plot advisory and agronomic outputs','MM-MD-0110 receives 22.4 pump hours and 1,120 cubic metres. Advice includes 100.2 t/ha yield, 1.55% five-day delay loss, fertigation doses, and low disease risk.'),
  ('rain','Figure 5  Rain-aware regional advice','With a 40 mm forecast, 894 irrigations receive the rain-skip status and none remain due today. These are model-based scenario counts, with estimated rather than measured savings.'),
 ],height=2.05)
 h(doc,'Interpretation of the displayed results')
 p(doc,'The plot advisory combines irrigation, stress, yield, delay loss, fertigation, and disease-risk outputs. Its explanatory text identifies lower soil moisture as the largest factor bringing irrigation forward. The map provides spatial context for those predictions. The rainfall comparison illustrates counterfactual advice under a changed forecast; it does not demonstrate a realised reduction in water extraction.',size=10.5,after=7)
 p(doc,'The saved advisory screenshot displays a target date of 30 September 2026. This is the interface date in the retained image, rather than an experimental observation date. A separate demonstration records the current local application.',size=9.5,after=0)

 page(doc,8,'8  Evaluation and Further Application Outputs')
 p(doc,'Table 2. Saved test results on generated labels for held-out plots. Regression scores use R-squared and MAE in the stated units; classification scores use accuracy and macro-F1. Values are read from models/metrics.json [10].',size=9.5,after=6)
 rows=[]
 short={'next_irrigation':'Irrigation timing (days)','water_requirement':'Water need (mm/day)','irrigation_depth':'Net depth (mm)','yield':'Yield (t/ha)','yield_loss':'Delay loss (percentage points)','water_stress':'Water stress','disease_risk':'Disease and pest risk'}
 for k in ['next_irrigation','water_requirement','irrigation_depth','yield','yield_loss','water_stress','disease_risk']:
  m=M['models'][k];v=m['test'];b=m['baseline'];reg='r2' in v
  rows.append([short[k],f"{v['r2']:.4f}" if reg else f"{100*v['accuracy']:.2f}%",f"{v['mae']:.4f}" if reg else f"{v['f1_macro']:.4f}",f"{b['r2']:.4f}" if reg else f"{100*b['accuracy']:.2f}%"])
 table(doc,['Target','R-squared / acc.','MAE / macro-F1','Baseline score'],rows,[2.65,1.4,1.4,1.4],9)
 p(doc,'All seven ensemble scores exceed their saved linear or logistic baseline. Timing MAE falls from 2.5796 to 0.6942 days. These comparisons assess approximation of the label generator and do not quantify field accuracy.',size=9.5,after=7)
 grid_panels(doc,[
  ('simulator','Figure 6  What-if simulation','The heat-wave preset shows severe water stress and recalculates irrigation need. Its delay curve and feature effects allow the user to inspect sensitivity to changed conditions.'),
  ('scheduler','Figure 7  Power-window scheduling','For Hosur, 20 pumps and a three-day horizon, the saved run completes 8 plots versus 5 under rotation. Estimated cane loss is 80.23 t versus 85.60 t.'),
  ('kannada','Figure 8  Kannada advisory','The same structured recommendation appears in Kannada. English, Hindi, and Marathi are also supported, with read-aloud dependent on browser voice availability.'),
  ('chat','Figure 9  Krishi Mitra conversation','English and Kannada questions retrieve the plot advisory. Keyword matching recognises irrigation intent and the plot ID, without an external language-model service.'),
 ],height=1.85)
 p(doc,'Model Insights also provides feature importance, predicted-versus-label plots, confusion matrices, and data-quality summaries. The 23 passing software tests include parity against saved Python fixtures and scheduling constraints. Neither these diagnostics nor the saved scenario comparisons constitute a field trial [10].',size=9.5,after=0)

 page(doc,9,'9  Conclusion and Future Scope')
 h(doc,'Conclusion')
 p(doc,'SMA implements a complete local software path from geofenced plot data to irrigation advice and pump scheduling. Its browser interface presents model predictions alongside hydraulic conversions, rainfall scenarios, stage-wise fertigation, language templates, and feature-sensitivity explanations. Python-to-JSON export allows the Node.js application to serve the trained tree ensembles without a separate Python prediction service [10].')
 p(doc,'The saved evaluation shows that the ensembles approximate the project\'s generated targets better than their linear or logistic baselines. Irrigation timing reaches an R-squared of 0.9646 with 0.6942 days MAE, while the classifiers obtain 87.83% water-stress accuracy and 87.29% disease-risk accuracy. Existing tests confirm the tested API, model-serving, language, and scheduling behaviour. The results demonstrate a functioning prototype and reproducible computational workflow [10].')
 h(doc,'Limitations')
 p(doc,'No independent field labels validate irrigation timing, yield, or disease risk. The input data are a static regional snapshot, crop age defaults to an estimate, and weather forecasts are manually selected scenarios. Soil hydraulics, radiation, equipment efficiency, and pump flow include simplifying assumptions. Tree spread is uncalibrated, and the scheduler\'s estimated benefit depends on the same generated yield-loss relationship that informs its priorities.')
 p(doc,'The interface can log accept or override feedback, but the training pipeline does not automatically incorporate those records. The local API also lacks a demonstrated production identity and access-control system. Wider deployment therefore requires additional agronomic, operational, and software engineering work [10].')
 h(doc,'Future scope')
 for t in [
  'Collect time-stamped irrigation events, measured discharge, soil profiles, disease observations, and harvest yield. Validate by season and location with preprocessing fitted only on training data.',
  'Replace estimated crop age with planting records and connect verified weather and earth-observation feeds. Maintain observation timestamps and validate the meaning of each input.',
  'Calibrate uncertainty against independent outcomes and evaluate error separately across taluks, growth stages, and extreme weather. Ask agronomists to review fertigation and risk rules.',
  'Compare the schedule with stronger optimisation methods using measured feeder capacity, pump discharge, interruptions, and farmer availability. Report actual water and energy use in a controlled pilot.',
  'Add role-based access, durable data storage, monitored services, and a reviewed feedback-to-retraining workflow. Test the four-language interaction with intended users.',
 ]:p(doc,t,'List Bullet',size=10.5,after=6)

 page(doc,10,'10  References')
 p(doc,'IEEE-style numbered references. Online sources checked on 19 October 2026. Project implementation, evaluation, and screenshot claims cite [10].',size=10,after=12)
 for ref in REFS:p(doc,ref,size=10,after=12)
 # Remove inherited theme overrides and title rules in the base DOCX template.
 for st in doc.styles:
  for node in st._element.xpath('.//w:pBdr'):
   node.getparent().remove(node)
  rf=st._element.find('.//'+qn('w:rFonts'))
  if rf is not None:
   for attr in list(rf.attrib):
    if 'Theme' in attr or 'theme' in attr:del rf.attrib[attr]
   rf.set(qn('w:ascii'),FONT);rf.set(qn('w:hAnsi'),FONT)
 for node in doc.settings._element.xpath('.//w:updateStylesOnOpen'):
  node.getparent().remove(node)
 doc.save(OUT/'SMA_Major_Project_Report_10_Pages.docx')

# Native DrawingML editing preserves template masters, theme, dimensions and branding.
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
REL='http://schemas.openxmlformats.org/package/2006/relationships'
def tag(s):a,b=s.split(':');return '{'+NS[a]+'}'+b
def el(s,**attrs):return E.Element(tag(s),{k:str(v) for k,v in attrs.items()})
def sub(par,s,**attrs):x=el(s,**attrs);par.append(x);return x
def emu(v):return str(round(v*914400))
class Slide:
 def __init__(self,pkg,source,k):
  self.pkg=pkg;self.k=k;self.id=1000;self.newrels=[]
  self.root=E.fromstring(pkg[f'ppt/slides/slide{source}.xml']);self.tree=self.root.find('p:cSld/p:spTree',NS)
  keep=[]
  for child in list(self.tree):
   if E.QName(child).localname in ['nvGrpSpPr','grpSpPr']:keep.append(child);continue
   if E.QName(child).localname=='pic':
    xf=child.find('p:spPr/a:xfrm/a:off',NS)
    if xf is not None and int(xf.get('y'))<100000:keep.append(child)
   elif child.find('p:spPr/a:xfrm/a:off',NS) is not None:
    off=child.find('p:spPr/a:xfrm/a:off',NS)
    if int(off.get('y'))>6400000 and not child.findall('.//a:t',NS):keep.append(child)
  for child in list(self.tree):
   if child not in keep:self.tree.remove(child)
  # Strip transitions and extensions from the source slide.
  for child in list(self.root):
   if E.QName(child).localname not in ['cSld','clrMapOvr']:self.root.remove(child)
  oldrels=E.fromstring(pkg[f'ppt/slides/_rels/slide{source}.xml.rels'])
  embeds=set(x.get(tag('r:embed')) for x in self.tree.findall('.//a:blip',NS))
  self.rels=E.Element('{'+REL+'}Relationships',nsmap={None:REL})
  for r in oldrels:
   if r.get('Id') in embeds or r.get('Type','').endswith('/slideLayout'):self.rels.append(deepcopy(r))
  self.notes=''
 def shape(self,x,y,w,h,fill=None,line=None,geometry='rect'):
  self.id+=1;s=el('p:sp');self.tree.append(s);nv=sub(s,'p:nvSpPr');sub(nv,'p:cNvPr',id=self.id,name=f'SMA shape {self.id}');sub(nv,'p:cNvSpPr',txBox=1);sub(nv,'p:nvPr')
  sp=sub(s,'p:spPr');xf=sub(sp,'a:xfrm');sub(xf,'a:off',x=emu(x),y=emu(y));sub(xf,'a:ext',cx=emu(w),cy=emu(h));geo=sub(sp,'a:prstGeom',prst=geometry);sub(geo,'a:avLst')
  if fill:sub(sub(sp,'a:solidFill'),'a:srgbClr',val=fill)
  else:sub(sp,'a:noFill')
  ln=sub(sp,'a:ln',w=12700)
  if line:sub(sub(ln,'a:solidFill'),'a:srgbClr',val=line)
  else:sub(ln,'a:noFill')
  return s
 def text(self,text,x,y,w,h,size=24,bold=False,color='000000',align='l',bullet=False,fill=None,line=None,pad=.02):
  s=self.shape(x,y,w,h,fill,line);body=sub(s,'p:txBody');bp=sub(body,'a:bodyPr',wrap='square',lIns=emu(pad),rIns=emu(pad),tIns=emu(pad),bIns=emu(pad),anchor='t');sub(bp,'a:noAutofit');sub(body,'a:lstStyle')
  for txt in (text if isinstance(text,list) else str(text).split('\n')):
   pa=sub(body,'a:p');pr=sub(pa,'a:pPr',algn=align,marL=emu(.22 if bullet else 0),indent=emu(-.2 if bullet else 0));sub(sub(pr,'a:lnSpc'),'a:spcPct',val=105000);sub(sub(pr,'a:spcAft'),'a:spcPts',val=650 if bullet else 0)
   if bullet:sub(pr,'a:buChar',char='•')
   else:sub(pr,'a:buNone')
   r=sub(pa,'a:r');rp=sub(r,'a:rPr',lang='en-IN',sz=int(size*100),b='1' if bold else '0');sub(sub(rp,'a:solidFill'),'a:srgbClr',val=color);sub(rp,'a:latin',typeface=FONT);sub(rp,'a:ea',typeface=FONT);sub(rp,'a:cs',typeface=FONT);sub(r,'a:t').text=txt;sub(pa,'a:endParaRPr',sz=int(size*100))
  return s
 def heading(self,t):self.text(t,1.52,1.55,11.4,.53,32,True,'C00000')
 def picture(self,name,x,y,w,h):
  self.id+=1;file=ASSETS/(name+'.png');im=Image.open(file);iw,ih=im.size;sc=min(w/iw,h/ih);ww,hh=iw*sc,ih*sc;x+=(w-ww)/2;y+=(h-hh)/2
  rid=f'rIdSMA{self.id}';target=f'../media/sma_{name}.png';r=E.SubElement(self.rels,'{'+REL+'}Relationship',Id=rid,Type=NS['r']+'/image',Target=target);self.pkg['ppt/media/sma_'+name+'.png']=file.read_bytes()
  s=sub(self.tree,'p:pic');nv=sub(s,'p:nvPicPr');sub(nv,'p:cNvPr',id=self.id,name='SMA '+name,descr='Actual application screenshot crop' if name!='architecture' else 'Architecture derived from source code');sub(nv,'p:cNvPicPr');sub(nv,'p:nvPr');bf=sub(s,'p:blipFill');bl=sub(bf,'a:blip');bl.set(tag('r:embed'),rid);sub(sub(bf,'a:stretch'),'a:fillRect');sp=sub(s,'p:spPr');xf=sub(sp,'a:xfrm');sub(xf,'a:off',x=emu(x),y=emu(y));sub(xf,'a:ext',cx=emu(ww),cy=emu(hh));sub(sub(sp,'a:prstGeom',prst='rect'),'a:avLst')
 def native_table(self,headers,rows,x,y,widths,rh=.62,size=17):
  self.id+=1;g=sub(self.tree,'p:graphicFrame');nv=sub(g,'p:nvGraphicFramePr');sub(nv,'p:cNvPr',id=self.id,name='SMA editable table');sub(nv,'p:cNvGraphicFramePr');sub(nv,'p:nvPr');xf=sub(g,'p:xfrm');sub(xf,'a:off',x=emu(x),y=emu(y));sub(xf,'a:ext',cx=emu(sum(widths)),cy=emu(rh*(1+len(rows))))
  dat=sub(sub(g,'a:graphic'),'a:graphicData',uri='http://schemas.openxmlformats.org/drawingml/2006/table');t=sub(dat,'a:tbl');sub(t,'a:tblPr',firstRow=1,bandRow=1);grid=sub(t,'a:tblGrid')
  for ww in widths:sub(grid,'a:gridCol',w=emu(ww))
  for ri,row in enumerate([headers]+rows):
   tr=sub(t,'a:tr',h=emu(rh))
   for text in row:
    tc=sub(tr,'a:tc');body=sub(tc,'a:txBody');sub(body,'a:bodyPr');sub(body,'a:lstStyle');pa=sub(body,'a:p');sub(pa,'a:pPr',algn='l');r=sub(pa,'a:r');rp=sub(r,'a:rPr',sz=int(size*100),b=1 if ri==0 else 0);sub(sub(rp,'a:solidFill'),'a:srgbClr',val='000000');sub(rp,'a:latin',typeface=FONT);sub(r,'a:t').text=str(text);sub(pa,'a:endParaRPr',sz=int(size*100));pr=sub(tc,'a:tcPr',marL=emu(.08),marR=emu(.08),marT=emu(.05),marB=emu(.05),anchor='ctr');sub(sub(pr,'a:solidFill'),'a:srgbClr',val='E8EDE8' if ri==0 else ('F5F8F3' if ri%2 else 'FFFFFF'))
    for edge in ['lnL','lnR','lnT','lnB']:sub(sub(sub(pr,'a:'+edge,w=6350),'a:solidFill'),'a:srgbClr',val='A8B4A8')
 def save(self,notes=''):
  # Native slide notes retain source citations and the interpretation of results.
  nt=E.Element(tag('p:notes'),nsmap=NS);cs=sub(nt,'p:cSld');tree=sub(cs,'p:spTree');nv=sub(tree,'p:nvGrpSpPr');sub(nv,'p:cNvPr',id=1,name='');sub(nv,'p:cNvGrpSpPr');sub(nv,'p:nvPr');sub(tree,'p:grpSpPr');s=sub(tree,'p:sp');nv=sub(s,'p:nvSpPr');sub(nv,'p:cNvPr',id=2,name='Notes');sub(nv,'p:cNvSpPr');nvpr=sub(nv,'p:nvPr');sub(nvpr,'p:ph',type='body',idx=1);sub(s,'p:spPr');body=sub(s,'p:txBody');sub(body,'a:bodyPr');sub(body,'a:lstStyle')
  for line in notes.split('\n'):pa=sub(body,'a:p');r=sub(pa,'a:r');sub(r,'a:t').text=line
  name=f'ppt/notesSlides/notesSlide{self.k}.xml';self.pkg[name]=E.tostring(nt,xml_declaration=True,encoding='UTF-8',standalone=True)
  nr=E.Element('{'+REL+'}Relationships',nsmap={None:REL});E.SubElement(nr,'{'+REL+'}Relationship',Id='rId1',Type=NS['r']+'/slide',Target=f'../slides/slide{self.k}.xml');self.pkg[f'ppt/notesSlides/_rels/notesSlide{self.k}.xml.rels']=E.tostring(nr,xml_declaration=True,encoding='UTF-8',standalone=True)
  E.SubElement(self.rels,'{'+REL+'}Relationship',Id='rIdSMANotes',Type=NS['r']+'/notesSlide',Target=f'../notesSlides/notesSlide{self.k}.xml')
  self.pkg[f'ppt/slides/slide{self.k}.xml']=E.tostring(self.root,xml_declaration=True,encoding='UTF-8',standalone=True);self.pkg[f'ppt/slides/_rels/slide{self.k}.xml.rels']=E.tostring(self.rels,xml_declaration=True,encoding='UTF-8',standalone=True)

def presentation():
 src=OUT/'dump/Original_LY_Major_Project_Template.pptx'
 pkg={i.filename:zipfile.ZipFile(src).read(i.filename) for i in zipfile.ZipFile(src).infolist()}
 # Avoid reading already-edited slides as a base.
 basepkg=deepcopy(pkg)
 def start(k,source=3):
  tmp=deepcopy(basepkg);s=Slide(tmp,source,k);s.pkg=pkg;return s
 s=start(1,1)
 s.text('Project Title',1.5,1.62,11.5,.6,32,True,align='ctr')
 s.text('Smart Moisture Advisor (SMA)',1.55,2.35,11.35,.6,30,True,'C00000',align='ctr')
 s.text('An AI-driven irrigation advisory system for sugarcane',1.55,3.02,11.35,.55,25,align='ctr')
 s.text('Members\nKartik Verma\nKushal Soni',1.55,4.02,11.35,1.25,25,True,align='ctr')
 s.text('Guide: Dr. Nilesh Yadav',1.55,5.6,11.35,.5,25,True,align='ctr')
 s.text('Department of Computer Engineering   2026-27',1.55,6.35,11.35,.35,19,align='ctr')
 s.save('Students: Kartik Verma and Kushal Soni. Guide: Dr. Nilesh Yadav. Institution and department follow the supplied template and report/build_report.js. Project title follows README.md and package.json. No other student or guide attribution is retained.')
 s=start(2);s.heading('Outline of Presentation');s.text(['Introduction','Summary / Finding of Literature Survey','System Design & Architecture','Proposed Methodology and Implementation','Application Outputs','Results','Conclusion and Future Scope','References'],1.65,2.25,11.1,4.4,25,bullet=True);s.save('This outline follows the major-project template structure and covers the implemented SMA application.')
 s=start(3);s.heading('Introduction');s.text([
 'SMA supports sugarcane irrigation decisions for 1,000 plots in Mandya, Karnataka.',
 'Plot vegetation, soil and climate inputs produce irrigation timing, water volume and pump duration.',
 'The interface links rain-aware advice with fertigation, risk estimates, a map and pump scheduling.',
 'Seven predictive models learn agronomy-generated labels. Field outcomes remain unverified.',
 'Advisories support English, Kannada, Hindi and Marathi, with accept or override feedback.',
 ],1.65,2.3,11.2,4.45,24,bullet=True);s.save('Sources: README.md, src/app.js, src/engine/advisory.js, src/engine/i18n.js, ml/train.py. Models estimate generated targets rather than measured field outcomes. Feedback is logged but not automatically used in training.')
 s=start(4);s.heading('Summary / Finding of Literature Survey');s.native_table(['Source','Approach or finding','Gap relative to SMA scope'],[
 ['Allen et al. [1]','FAO crop water balance and ET','Requires suitable local inputs and calibration'],
 ['Doorenbos and Kassam [2]','Yield response to water deficit','Does not provide an integrated plot interface'],
 ['Kamienski et al. [3]','SWAMP IoT platform and pilots','Deployment and scaling require infrastructure'],
 ['Breiman [4] and Friedman [5]','Random Forest and Gradient Boosting','Require representative targets and evaluation'],
 ['Current project [10]','Integrates advice, maps and scheduling','Independent field validation is still needed'],
 ],1.57,2.25,[2.7,4.05,4.55],.73,18)
 s.text('The comparison identifies project requirements, without claiming superiority over the cited systems.',1.6,6.75,11.2,.3,15);s.save('\n'.join(REFS[:6])+'\nComparative gaps are scoped to this implementation. No head-to-head evaluation with SWAMP or other irrigation platforms was performed.')
 s=start(5);s.heading('System Design & Architecture');s.picture('architecture',1.55,2.2,11.35,4.58);s.save('Architecture derived from ml/train.py, src/data/farms.js, src/ml/registry.js, src/ml/treeEnsemble.js, src/app.js and src/engine/. Offline training exports JSON. Runtime prediction occurs locally in Node.js. Feedback is appended to JSONL, not fed automatically into training.')
 s=start(6);s.heading('Proposed Methodology and Implementation')
 s.text(['Clean plot data and create twelve scenarios per plot.',
 'Generate noisy FAO-based labels for irrigation, crop demand and risk.',
 'Split by plot: 800 training plots and 200 held-out plots.',
 'Train ensembles, compare linear or logistic baselines, then export JSON.',
 'Serve predictions through Express and apply hydraulics, fertigation and scheduling rules.'],1.65,2.25,6.45,4.55,22,bullet=True)
 stages=['1,000 plot records','12,000 generated scenarios','Grouped train / test evaluation','JSON model export','JavaScript advisory and scheduler']
 for i,t in enumerate(stages):
  y=2.28+i*.86;s.text(t,8.65,y,4.15,.58,19,True,align='ctr',fill='F0F4EF',line='70846F',pad=.07)
  if i<4:s.shape(10.55,y+.6,.3,.21,fill='70846F',geometry='downArrow')
 s.save('Sources: ml/train.py and ml/agronomy.py. Label generation is supervised learning from a simulated agronomic teacher, not a physics-informed neural network. NDVI imputation currently precedes the grouped split. WSPT is a heuristic for parallel lanes and power windows, not a proof of optimal constrained scheduling.')
 s=start(7);s.heading('Application Outputs')
 # Three readable visuals rather than full-page application thumbnails.
 s.picture('advisor',1.6,2.25,5.45,2.75);s.text('Plot advice, yield, fertigation and risk',1.6,5.05,5.45,.4,18,True)
 s.picture('map',7.45,2.25,5.2,2.75);s.text('Geofenced plots coloured by stress',7.45,5.05,5.2,.4,18,True)
 s.text(['What-if Simulator recomputes outputs when field conditions change.',
 'Krishi Mitra uses keyword intent detection in four supported languages.',
 'Farmers and supervisors can accept or override the recommendation.'],1.65,5.63,11.05,1.25,20,bullet=True)
 s.save('Actual screenshot crops from presentation/screenshots/03_advisor.png and 02b_map_zoom.png. Example MM-MD-0110 is in the summer dry-spell scenario: 22.4 hours pumping, 1,120 cubic metres, 100.2 t/ha predicted yield, 1.55% predicted loss after a five-day delay. These values are scenario outputs, not measured farm outcomes. Capture date in the original screenshot: 30 September 2026. Chatbot: src/engine/chat.js.')
 s=start(8);s.heading('Results')
 rows=[]
 for k,label in [('next_irrigation','Irrigation timing'),('water_requirement','Crop water requirement'),('irrigation_depth','Irrigation depth'),('yield','Cane yield'),('yield_loss','Delay loss'),('water_stress','Water stress'),('disease_risk','Disease risk')]:
  v=M['models'][k]['test'];b=M['models'][k]['baseline'];reg='r2'in v
  rows.append([label,f"{v['r2']:.4f}" if reg else f"{v['accuracy']*100:.2f}%",f"{b['r2']:.4f}" if reg else f"{b['accuracy']*100:.2f}%"])
 s.native_table(['Target','Ensemble score','Baseline score'],rows,1.57,2.25,[3.85,2.2,2.2],.43,17)
 s.text('Hosur schedule',10.03,2.55,2.7,.5,21,True)
 s.text('5.37 t',10.03,3.35,2.7,.6,32,True,'C00000')
 s.text('less estimated cane loss\n8 plots completed vs 5\n20 pumps over 3 days',10.03,4.15,2.7,1.45,18)
 s.text('R-squared for regression, accuracy for classifiers. Timing MAE: 0.6942 days.\nEvaluation uses generated labels on held-out plots. Savings and cane loss are scenario estimates.',1.65,6.15,11, .75,19)
 s.save('Source: models/metrics.json. Regression baselines are Linear Regression, classifiers use Logistic Regression. Timing baseline MAE 2.5796 days, ensemble 0.6942 days. Scheduler: report/data/results.json, schedule key, Hosur, 20 pumps, windows 06:00-10:00 and 17:00-20:00, three days. Optimised estimated cane loss 80.23 t, FCFS 85.60 t. Software tests: 23 passed on 19 October 2026. Results assess generated labels and scenario simulations, with no field trial.')
 s=start(9);s.heading('Conclusion and Future Scope');s.text([
 'SMA integrates plot-level prediction, irrigation advice and power-window scheduling in a working web prototype.',
 'JSON model export supports local JavaScript inference. All 23 existing software tests pass.',
 'The saved metrics show agreement with generated agronomic labels. Independent field validation is the main next step.',
 'Future work: measured irrigation and yield records, verified weather feeds, actual planting dates and calibrated uncertainty.',
 'Operational development: agronomist review, field pilots, role-based access and a reviewed retraining workflow.',
 ],1.65,2.3,11.2,4.45,24,bullet=True);s.save('Sources: current source code, models/metrics.json, report/data/extra.json and execution of node --test test/api.test.js test/engine.test.js test/models.test.js. Live feed connectors, authenticated production operation and automatic feedback retraining are proposed work.')
 s=start(10);s.heading('References')
 short=[
 '[1] Allen et al., Crop Evapotranspiration, FAO Paper 56, 1998. fao.org/4/x0490e/x0490e00.htm',
 '[2] Doorenbos and Kassam, Yield Response to Water, FAO Paper 33, 1979.',
 '[3] Kamienski et al., Smart Water Management Platform, Sensors 19(2), 276, 2019. doi:10.3390/s19020276',
 '[4] Breiman, Random forests, Machine Learning 45, 5-32, 2001. doi:10.1023/A:1010933404324',
 '[5] Friedman, A gradient boosting machine, Annals of Statistics 29(5), 1189-1232, 2001. doi:10.1214/aos/1013203451',
 '[6] Pedregosa et al., Scikit-learn, JMLR 12, 2825-2830, 2011. jmlr.org/papers/v12/pedregosa11a.html',
 '[7] Scikit-learn developers, GroupShuffleSplit API. scikit-learn.org/stable/modules/generated/sklearn.model_selection.GroupShuffleSplit.html',
 '[8] Express contributors, Express 5.x API. expressjs.com/en/5x/api/',
 '[9] Leaflet contributors, Leaflet API reference. leafletjs.com/reference.html',
 '[10] Verma and Soni, Smart Moisture Advisor source and saved outputs. github.com/Interior-Gardener/SMA',
 ]
 s.text(short,1.62,2.25,11.1,4.65,18,bullet=False);s.save('\n'.join(REFS))
 s=start(11);s.text('Thank You',1.6,2.9,11.1,.85,44,True,'C00000',align='ctr');s.text('Smart Moisture Advisor',1.6,4.1,11.1,.55,30,True,align='ctr');s.text('Kartik Verma and Kushal Soni\nProject Guide: Dr. Nilesh Yadav',1.6,5.15,11.1,1.0,25,align='ctr');s.save('Students: Kartik Verma and Kushal Soni. Project Guide: Dr. Nilesh Yadav.')
 # Ensure content types for notes and PNG. Remove all unreachable old-project parts.
 ct=E.fromstring(pkg['[Content_Types].xml']);ctns='{http://schemas.openxmlformats.org/package/2006/content-types}'
 for i in range(1,12):
  part=f'/ppt/notesSlides/notesSlide{i}.xml'
  if not any(e.get('PartName')==part for e in ct):E.SubElement(ct,ctns+'Override',PartName=part,ContentType='application/vnd.openxmlformats-officedocument.presentationml.notesSlide+xml')
 if not any(e.get('Extension')=='png' for e in ct):E.SubElement(ct,ctns+'Default',Extension='png',ContentType='image/png')
 # Metadata and thumbnail should not retain the old title or old chemistry image.
 for part in ['docProps/core.xml']:
  if part in pkg:
   core=E.fromstring(pkg[part]);dc={'dc':'http://purl.org/dc/elements/1.1/','cp':'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'}
   for field,val in [('dc:title','Smart Moisture Advisor'),('dc:creator','Kartik Verma and Kushal Soni'),('dc:subject','Major project presentation')]:
    e=core.find(field,dc)
    if e is None:e=E.SubElement(core,'{'+dc[field.split(':')[0]]+'}'+field.split(':')[1])
    e.text=val
   pkg[part]=E.tostring(core,xml_declaration=True,encoding='UTF-8',standalone=True)
 rootrels=E.fromstring(pkg['_rels/.rels'])
 for r in list(rootrels):
  if r.get('Type','').endswith('/thumbnail'):rootrels.remove(r)
 pkg['_rels/.rels']=E.tostring(rootrels,xml_declaration=True,encoding='UTF-8',standalone=True)
 reachable={'[Content_Types].xml','_rels/.rels'};queue=['']
 while queue:
  owner=queue.pop();relname='_rels/.rels' if not owner else posixpath.join(posixpath.dirname(owner),'_rels',posixpath.basename(owner)+'.rels')
  if relname not in pkg:continue
  reachable.add(relname)
  for r in E.fromstring(pkg[relname]):
   if r.get('TargetMode')=='External':continue
   target=r.get('Target','');part=posixpath.normpath(posixpath.join(posixpath.dirname(owner),target)).lstrip('/')
   if target.startswith('/'):part=target[1:]
   if part in pkg and part not in reachable:reachable.add(part);queue.append(part)
 for e in list(ct):
  if e.tag.endswith('Override') and e.get('PartName','').lstrip('/') not in reachable:ct.remove(e)
 pkg['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
 target=OUT/'SMA_Major_Project_Presentation.pptx'
 with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
  for n in sorted(reachable):z.writestr(n,pkg[n])
 (BUILD/'template_fidelity.json').write_text(json.dumps({'source':str(src),'sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'slide_size_emu':[12192000,6858000],'slides':11,'preserved':'Original theme, masters, slide layouts, red border graphics, institution logo and Times New Roman typography','changes':'Replaced project content, diagram, screenshot assets, notes and chemistry-themed closing image. Video section replaced by application outputs.','authoring':'Native DrawingML edits, rendered in installed PowerPoint. Artifact Tool runtime and dependency-loader tool are not exposed in this Windows session.'},indent=2))

if __name__=='__main__':
 import sys
 architecture();crops();report()
 if '--report-only' not in sys.argv:presentation()
 print('CREATED report DOCX'+('' if '--report-only' in sys.argv else ' and template-based PPTX')+' in',OUT)
